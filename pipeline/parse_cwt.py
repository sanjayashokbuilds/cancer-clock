"""Unify newer combined CSVs and older provider workbooks without double counting."""
from __future__ import annotations

import logging
from pathlib import Path
import re

import numpy as np
from openpyxl import load_workbook
import pandas as pd

from config import ROOT_DIR
from pipeline.common import DataError, assert_unique, metadata, normal, numeric, read_csv, require_columns

LOG = logging.getLogger(__name__)
DIMS = ['Basis', 'Org_Code', 'Parent_Org', 'Org_Name', 'Standard_or_Item', 'Cancer_Type',
        'Referral_Route_or_Stage', 'Treatment_Modality']
BANDS = ['Within_14_days', 'In_15_to_28_days', 'In_29_to_42_days', 'In_43_to_62_days',
         'After_62_days', 'Within_31_days', 'In_32_to_38_days', 'In_39_to_48_days',
         'In_49_to_62_days', 'In_63_to_76_days', 'In_77_to_90_days', 'In_91_to_104_days',
         'After_104_days', 'In_15_to_16_days', 'In_17_to_21_days', 'In_22_to_28_days', 'After_28_days']
COUNTS = ['Total', 'Within', 'After'] + BANDS
KEY = ['Org_Code', 'period', 'Standard_or_Item', 'Cancer_Type', 'Referral_Route_or_Stage', 'Treatment_Modality']
PROVENANCE = ['published_on', 'publication_date_basis', 'source_url', 'source_sha256', 'source_version', 'historical_vintage_verified']


def legacy_header(values: list[object], standard: str) -> dict[int, str]:
    """Translate published header labels, keeping duplicate Within fields distinct."""
    mapping = {}
    total_index = next((i for i, v in enumerate(values) if normal(v).upper() == 'TOTAL'), None)
    if total_index is None:
        raise DataError('Legacy CWT table has no TOTAL column')
    for i, value in enumerate(values):
        text = re.sub(r'\s*\(\d+\)', '', normal(value)).upper()
        if text.startswith('ODS CODE'):
            mapping[i] = 'Org_Code'
        elif text == 'ACCOUNTABLE PROVIDER':
            mapping[i] = 'Org_Name'
        elif text in {'REFERRAL ROUTE', 'TREATMENT STAGE'}:
            mapping[i] = 'Referral_Route_or_Stage'
        elif text in {'CANCER TYPE', 'SUSPECTED CANCER OR BREAST SYMPTOMATIC', 'SUSPECTED TYPE OF CANCER', 'REFERRAL REASON'}:
            mapping[i] = 'Cancer_Type'
        elif text == 'TREATMENT MODALITY':
            mapping[i] = 'Treatment_Modality'
        elif text == 'TOTAL':
            mapping[i] = 'Total'
        elif text.startswith(('TREATED WITHIN', 'TOLD WITHIN')):
            mapping[i] = 'Performance'
        elif i == total_index + 1 and standard in {'FDS', '31D', '62D'}:
            mapping[i] = 'Within'
        elif i == total_index + 2 and standard in {'FDS', '31D', '62D'}:
            mapping[i] = 'After'
        else:
            match = re.fullmatch(r'(WITHIN|AFTER) (\d+) DAYS', text)
            interval = re.fullmatch(r'IN (\d+) TO (\d+) DAYS', text)
            if match:
                mapping[i] = f'{match[1].title()}_{match[2]}_days'
            elif interval:
                mapping[i] = f'In_{interval[1]}_to_{interval[2]}_days'
    return mapping


def read_legacy(path: Path) -> pd.DataFrame:
    """Read actual provider tables; skip title rows, repeated headers and footnotes."""
    rows = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            # NHS exports can incorrectly declare the used range as A1 only.
            sheet.reset_dimensions()
            title = sheet.title.upper()
            if title.startswith('28-DAY'):
                standard = 'FDS'
            elif title.startswith('31-DAY'):
                standard = '31D'
            elif title.startswith('62-DAY'):
                standard = '62D'
            elif title.startswith('USC REFERRAL'):
                standard = 'Urgent Suspected Cancer referral'
            elif title.startswith('UBS REFERRAL'):
                standard = 'Urgent Breast Symptomatic referral'
            else:
                continue
            mapping = None
            for values in sheet.iter_rows(values_only=True):
                if any(str(v).upper().startswith('ODS CODE') for v in values if v is not None):
                    mapping = legacy_header(list(values), standard)
                    continue
                if mapping is None:
                    continue
                row = {column: values[i] if i < len(values) else None for i, column in mapping.items()}
                code, name = normal(row.get('Org_Code') or ''), normal(row.get('Org_Name') or '')
                if name == 'ALL ENGLISH PROVIDERS':
                    code = 'Total'
                if not re.fullmatch(r'[A-Z0-9]{3,5}|Total', code) or not name:
                    continue
                row.update(Basis='Provider', Org_Code=code, Parent_Org='', Standard_or_Item=standard)
                row.setdefault('Cancer_Type', 'ALL CANCERS')
                row.setdefault('Referral_Route_or_Stage', '')
                row.setdefault('Treatment_Modality', 'ALL MODALITIES' if standard in {'31D', '62D'} else '')
                rows.append(row)
    finally:
        workbook.close()
    if not rows:
        raise DataError(f'No recognised provider tables in {path.name}')
    return pd.DataFrame(rows)


def clean(frame: pd.DataFrame, item: dict[str, str]) -> pd.DataFrame:
    """Validate counts and preserve fractional shared treatment counts."""
    require_columns(frame, DIMS + ['Total'], f'CWT {item["period"]}')
    frame = frame.copy()
    for column in DIMS:
        frame[column] = frame[column].fillna('').map(normal)
    frame = frame.loc[frame.Basis.str.casefold().eq('provider')].copy()
    frame['Referral_Route_or_Stage'] = frame.Referral_Route_or_Stage.str.upper().replace({'SUBSEQUENT TREATMENT': 'SUBSEQUENT'})
    frame['Treatment_Modality'] = frame.Treatment_Modality.str.upper()
    for column in COUNTS:
        frame[column] = numeric(frame[column] if column in frame else pd.Series(np.nan, index=frame.index), column)
    # Retain the source percentage as text and calculate the canonical proportion.
    frame['Performance_source'] = frame.get('Performance', pd.Series('', index=frame.index)).fillna('').astype(str)
    frame['Performance'] = frame.Within.div(frame.Total.where(frame.Total.gt(0)))
    standards = frame.Standard_or_Item.isin(['FDS', '31D', '62D'])
    complete = standards & frame[['Total', 'Within', 'After']].notna().all(axis=1)
    if not np.allclose(frame.loc[complete, 'Total'], frame.loc[complete, 'Within'] + frame.loc[complete, 'After'], atol=1e-6):
        raise DataError(f'CWT {item["period"]}: Total != Within + After')
    if frame.loc[standards, 'Within'].gt(frame.loc[standards, 'Total']).any():
        raise DataError('CWT within count exceeds total')
    frame = metadata(frame, item)
    frame['source_layout'] = 'legacy_provider_workbook' if item['format'] == 'xlsx' else 'combined_csv'
    assert_unique(frame, KEY, 'CWT')
    return frame


def tidy(frame: pd.DataFrame) -> pd.DataFrame:
    """One row per provider/month/source cancer label, with explicit metric scopes."""
    index = ['Org_Code', 'period', 'Cancer_Type']
    pieces = []
    definitions = [
        ('fds', 'FDS', 'ALL ROUTES', ''),
        ('d62_usc', '62D', 'URGENT SUSPECTED CANCER', 'ALL MODALITIES'),
        ('d62_all', '62D', 'ALL ROUTES', 'ALL MODALITIES'),
        ('d31_all', '31D', 'ALL STAGES', 'ALL MODALITIES'),
        ('d31_surgery', '31D', 'ALL STAGES', 'SURGERY'),
        ('d31_radiotherapy', '31D', 'ALL STAGES', 'RADIOTHERAPY'),
        ('d31_drugs', '31D', 'ALL STAGES', 'ANTI-CANCER DRUG REGIMEN'),
        ('usc_first_seen', 'Urgent Suspected Cancer referral', '', ''),
    ]
    for prefix, standard, route, modality in definitions:
        rows = frame.loc[frame.Standard_or_Item.eq(standard) & frame.Referral_Route_or_Stage.eq(route) & frame.Treatment_Modality.eq(modality)]
        if rows.empty:
            continue
        metrics = ['Total', 'Within', 'After', 'Performance']
        if standard == '62D':
            metrics += ['Within_31_days', 'In_32_to_38_days', 'In_39_to_48_days', 'In_49_to_62_days', 'In_63_to_76_days', 'In_77_to_90_days', 'In_91_to_104_days', 'After_104_days']
        if standard.endswith('referral'):
            metrics = ['Total']
        assert_unique(rows, index, prefix)
        pieces.append(rows.set_index(index)[metrics].rename(columns=lambda c: prefix + '_' + c.lower()))
    result = pd.concat(pieces, axis=1).reset_index()
    # Source cancer labels intentionally stay distinct (suspected vs confirmed).
    lookup = frame[['Org_Code', 'period', 'Org_Name'] + PROVENANCE].drop_duplicates()
    assert_unique(lookup, ['Org_Code', 'period'], 'CWT metadata')
    result = result.merge(lookup, on=['Org_Code', 'period'], validate='many_to_one')
    return result.rename(columns={'Org_Code': 'org_code', 'Org_Name': 'org_name', 'Cancer_Type': 'cancer_type'})


def first_seen(frame: pd.DataFrame) -> pd.DataFrame:
    """Sum mutually exclusive USC categories only when no ALL CANCERS total exists."""
    rows = frame.loc[frame.Standard_or_Item.eq('Urgent Suspected Cancer referral')].copy()
    result = []
    for (code, period), group in rows.groupby(['Org_Code', 'period']):
        overall = group.loc[group.Cancer_Type.eq('ALL CANCERS')]
        total = overall.Total.iloc[0] if len(overall) == 1 else group.Total.sum(min_count=len(group))
        record = {'org_code': code, 'period': period, 'usc_first_seen': total,
                  'org_name': group.Org_Name.iloc[0], 'measure': 'completed_first_outpatient_attendances_not_referrals_received'}
        record.update({column: group[column].iloc[0] for column in PROVENANCE})
        result.append(record)
    return pd.DataFrame(result)


def parse_all(manifest: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Build canonical CWT, separate national totals, tidy metrics and volume proxy."""
    parts = []
    for item in manifest.loc[manifest.dataset.eq('cwt')].to_dict('records'):
        LOG.info('Parsing CWT %s (%s)', item['period'], item['format'])
        path = ROOT_DIR / item['local_path']
        frame = read_legacy(path) if item['format'] == 'xlsx' else read_csv(path)
        parts.append(clean(frame, item))
    full = pd.concat(parts, ignore_index=True)
    national = full.loc[full.Org_Code.str.casefold().eq('total')].copy()
    providers = full.loc[~full.Org_Code.str.casefold().eq('total')].copy()
    assert_unique(providers, KEY, 'CWT all months')
    return {'cwt': providers, 'cwt_national': national, 'cwt_trust_month': tidy(providers),
            'cwt_first_seen': first_seen(providers)}
