"""Small synthetic fixtures test parser failures; they are never demo data."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import Workbook

from pipeline.common import DataError, numeric, period_from_text
from pipeline.fetch import candidate, select_sources, Downloader
from pipeline.parse_cwt import clean, first_seen, read_legacy, tidy, DIMS
from pipeline.parse_dm01 import parse_zip
from pipeline.parse_backlog import parse_workbook
from pipeline.vintages import available_at


@pytest.fixture
def item() -> dict:
    """Source metadata for synthetic tests, deliberately marked as unverified."""
    return dict(period='2026-08-01', published_on='2026-10-14', publication_date_basis='test_assumption',
                url='https://www.england.nhs.uk/test.csv', sha256='synthetic-test', version='provisional', format='csv')


def cwt_row(**overrides: object) -> dict:
    """One synthetic, internally consistent standard row."""
    row = dict.fromkeys(DIMS, '')
    row.update(Basis='Provider', Org_Code='RTH', Org_Name='Test provider', Standard_or_Item='62D',
               Cancer_Type='ALL CANCERS', Referral_Route_or_Stage='Urgent Suspected Cancer',
               Treatment_Modality='ALL MODALITIES', Total='10.5', Within='8', After='2.5')
    row.update(overrides)
    return row


@pytest.mark.parametrize('text,expected', [('January 2023 – 24 (Final)', '2024-01-01'),
    ('October 2023-24 (Final)', '2023-10-01'), ('DM01-AUGUST-2026-full-extract', '2026-08-01')])
def test_period_parsing(text: str, expected: str) -> None:
    """Fiscal page labels must not move January backwards a year."""
    assert period_from_text(text) == pd.Timestamp(expected)


def test_source_preference() -> None:
    """Select Final without inadvertently selecting a duplicate ICB version."""
    base = 'https://www.england.nhs.uk/statistics/wp-content/uploads/sites/2/2026/10/'
    values = [candidate('cwt', 'August 2026 Combined CSV ' + name, base + str(i) + '.csv', base)
              for i, name in enumerate(['Provisional', 'Final', 'Final New ICB Structure'])]
    assert select_sources(values)[0]['url'].endswith('1.csv')


def test_cached_download_makes_no_request(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """A completed cached source is reused on the next download."""
    import hashlib
    url = 'https://www.england.nhs.uk/example.csv'
    target = tmp_path / 'cwt' / (hashlib.sha256(url.encode()).hexdigest()[:20] + '.csv')
    target.parent.mkdir(); target.write_text('example')
    client = Downloader(tmp_path)
    def fail(*args: object, **kwargs: object) -> None:
        raise AssertionError('Network should not be called')
    monkeypatch.setattr(client.session, 'get', fail)
    assert client.get(url, 'cwt') == target


def test_missing_values_are_not_zero() -> None:
    """Suppressed cells stay missing and shared half-patient counts stay fractional."""
    values = numeric(pd.Series(['*', '', '1,200', '2.5']), 'test')
    assert values.iloc[:2].isna().all()
    assert values.iloc[2:].tolist() == [1200, 2.5]
    with pytest.raises(DataError):
        numeric(pd.Series(['unexpected']), 'test')


def test_cwt_count_checks_and_duplicates(item: dict) -> None:
    """Do not accept inconsistent counts or double-count duplicate provider rows."""
    cleaned = clean(pd.DataFrame([cwt_row()]), item)
    assert cleaned.Total.iloc[0] == 10.5
    assert tidy(cleaned).d62_usc_total.iloc[0] == 10.5
    with pytest.raises(DataError, match='Total !='):
        clean(pd.DataFrame([cwt_row(Total='12')]), item)
    with pytest.raises(DataError, match='duplicate'):
        clean(pd.DataFrame([cwt_row(), cwt_row()]), item)


def test_first_seen_avoids_overlapping_total(item: dict) -> None:
    """The overall published row takes precedence over its component categories."""
    rows = [cwt_row(Standard_or_Item='Urgent Suspected Cancer referral', Cancer_Type=cancer,
                    Referral_Route_or_Stage='', Treatment_Modality='', Total=total)
            for cancer, total in [('ALL CANCERS', '10'), ('Suspected lung cancer', '4'), ('Suspected skin cancer', '6')]]
    assert first_seen(clean(pd.DataFrame(rows), item)).usc_first_seen.iloc[0] == 10


def test_legacy_sheet_headers_and_national_total(tmp_path: Path, item: dict) -> None:
    """Older workbooks include national rows with blank ODS code and repeated headers."""
    workbook = Workbook(); sheet = workbook.active; sheet.title = '62-DAY (ALL CANCER)'
    header = ['', 'ODS CODE (1)', 'ACCOUNTABLE PROVIDER', 'REFERRAL ROUTE', 'CANCER TYPE', 'TOTAL', 'WITHIN 62 DAYS', 'AFTER 62 DAYS', 'TREATED WITHIN 62 DAYS']
    sheet.append(['Title']); sheet.append(header)
    sheet.append(['', 'RTH', 'Test provider', 'Urgent Suspected Cancer', 'ALL CANCERS', 10.5, 8, 2.5, 8/10.5])
    sheet.append(header)
    sheet.append(['', '', 'ALL ENGLISH PROVIDERS', 'Urgent Suspected Cancer', 'ALL CANCERS', 10.5, 8, 2.5, 8/10.5])
    path = tmp_path / 'legacy.xlsx'; workbook.save(path)
    result = clean(read_legacy(path), {**item, 'format': 'xlsx'})
    assert set(result.Org_Code) == {'RTH', 'Total'}


def test_vintage_filter_requires_explicit_approximation() -> None:
    """Revised historical values must never silently pass as verified vintages."""
    frame = pd.DataFrame({'published_on': pd.to_datetime(['2026-10-14', '2026-11-14']), 'historical_vintage_verified': [False, False]})
    with pytest.raises(DataError, match='unverified'):
        available_at(frame, '2026-10-24')
    assert len(available_at(frame, '2026-10-24', allow_approximate=True)) == 1
    assert available_at(frame, '2026-09-01').empty


def make_dm_zip(path: Path, suppressed: bool = False, duplicate: bool = False) -> None:
    """Build a tiny DM01 input with all seven recognised test labels."""
    from pipeline.parse_dm01 import TESTS
    records = []
    for test in TESTS:
        for commissioner in ('A', 'B'):
            row = {'Period': 'DM01-AUGUST-2026', 'Provider Org Code': 'RTH', 'Provider Org Name': 'Test',
                   'Commissioner Org Code': commissioner, 'Diagnostic Tests': test, 'Total WL': '14',
                   'Waiting List Activity': '1', 'Planned Activity': '2', 'Unscheduled Activity': '3', 'Total Activity': '6'}
            for i in range(13):
                row[f'{i:02d} < {i+1:02d} ' + ('Week' if i == 0 else 'Weeks')] = '1'
            row['13+ Weeks'] = '*' if suppressed and commissioner == 'A' else '1'
            records.append(row)
    if duplicate:
        records.append(records[0])
    with ZipFile(path, 'w') as archive:
        archive.writestr('extract.csv', pd.DataFrame(records).to_csv(index=False))


def test_dm01_aggregation_and_suppression(tmp_path: Path, item: dict) -> None:
    """Sum commissioners once; any missing age cell makes the long-wait count unknown."""
    path = tmp_path / 'dm.zip'; make_dm_zip(path)
    frame = parse_zip(path, item, {'RTH'})
    assert len(frame) == 7
    assert frame.total_waiting_list.eq(28).all()
    assert frame.waiting_6plus.eq(16).all()
    assert frame.activity.eq(12).all()
    make_dm_zip(path, suppressed=True)
    assert parse_zip(path, item, {'RTH'}).waiting_6plus.isna().all()
    make_dm_zip(path, duplicate=True)
    with pytest.raises(DataError, match='duplicate'):
        parse_zip(path, item, {'RTH'})


def test_dm01_stage4_excludes_nonc_without_changing_legacy_model_inputs(tmp_path: Path, item: dict) -> None:
    """English reference removes NONC; Stage 3's original parser default is preserved."""
    path=tmp_path/'dm.zip';make_dm_zip(path)
    with ZipFile(path) as archive:
        rows=pd.read_csv(BytesIO(archive.read('extract.csv')))
    rows.loc[rows['Commissioner Org Code'].eq('B'),'Commissioner Org Code']='NONC'
    with ZipFile(path,'w') as archive:
        archive.writestr('extract.csv',rows.to_csv(index=False))
    assert parse_zip(path,item,{'RTH'}).total_waiting_list.eq(28).all()
    english=parse_zip(path,item,None,exclude_non_english=True)
    assert english.total_waiting_list.eq(14).all()
    assert english.waiting_6plus.eq(8).all()


def test_backlog_missing_week_and_dtt_reconciliation(tmp_path: Path, item: dict) -> None:
    """Retain missing provider counts; validate the national diagnosis/treatment split."""
    workbook = Workbook(); national = workbook.active; national.title = '62 day PTL'
    national.append(['title']); national.append(['', '', 'Total', 'No DTT', 'With DTT'])
    national.append(['', 'w-e 02 Aug 26', 10, 7, 3]); national.append(['', 'w-e 09 Aug 26', 12, 8, 4])
    provider = workbook.create_sheet('62 Day PTL by Provider')
    provider.append(['title']); provider.append(['ORG CODE', 'ORG NAME', 'w-e 02 Aug 26', 'w-e 09 Aug 26'])
    provider.append(['RTH', 'Test provider', 10, None])
    path = tmp_path / 'backlog.xlsx'; workbook.save(path)
    data = parse_workbook(path, item, start='2026-08-01')
    assert data['backlog_trust'].backlog.isna().sum() == 1
    assert data['backlog_trust'].published_on.eq(pd.Timestamp('2026-10-14')).all()
    national.cell(3, 3, 99); workbook.save(path)
    with pytest.raises(DataError, match='reconcile'):
        parse_workbook(path, item)


def test_dm01_source_label_typo_and_cdc_exclusion() -> None:
    """The real March 2025 page has a wrong year in its link label and a CDC-only alternative."""
    from pipeline.fetch import is_dm01_extract
    url = 'https://www.england.nhs.uk/statistics/wp-content/uploads/sites/2/2026/04/DM01-MARCH-2025-full-extract_M0N7G.zip'
    title = 'CSV Extract – March 2026 – All Provider – Commissioner Data (Revised 16.04.2026)'
    assert is_dm01_extract(title, url)
    assert candidate('dm01', title, url, 'source')['period'] == '2025-03-01'
    assert not is_dm01_extract('CSV Extract – March 2025 – CDC – All Provider – Commissioner Data', url.replace('M0N7G', 'CDC'))
